"""本地知识库处理与 LangChain 问答服务。"""
import json
import hashlib
import logging
import math
import shutil
import time
import uuid
from pathlib import Path
from threading import RLock

from fastapi import HTTPException, UploadFile
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_core.prompts import ChatPromptTemplate
from langchain_core.embeddings import Embeddings
from langchain_openai import ChatOpenAI
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader
from docx import Document as WordDocument

from .config import settings

logger = logging.getLogger(__name__)
SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt"}
_embeddings = None
_vector_store = None
_load_lock = RLock()


class LocalChineseHashEmbeddings(Embeddings):
    """Small offline character n-gram embeddings for reliable local Chinese retrieval."""
    dimensions = 1024

    def _embed(self, text: str) -> list[float]:
        vector = [0.0] * self.dimensions
        normalized = "".join(text.lower().split())
        features = []
        for size in (1, 2, 3):
            features.extend(normalized[index:index + size] for index in range(max(0, len(normalized) - size + 1)))
        for feature in features:
            digest = hashlib.blake2b(feature.encode("utf-8"), digest_size=4).digest()
            bucket = int.from_bytes(digest, "big") % self.dimensions
            vector[bucket] += 1.0 if digest[0] % 2 else -1.0
        length = math.sqrt(sum(value * value for value in vector))
        return [value / length for value in vector] if length else vector

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [self._embed(text) for text in texts]

    def embed_query(self, text: str) -> list[float]:
        return self._embed(text)


def _get_embeddings():
    global _embeddings
    with _load_lock:
        if _embeddings is None:
            _embeddings = LocalChineseHashEmbeddings()
    return _embeddings


def _get_store():
    global _vector_store
    with _load_lock:
        if _vector_store is None:
            _vector_store = Chroma(
                collection_name="enterprise_knowledge",
                persist_directory=str(settings.chroma_dir),
                embedding_function=_get_embeddings(),
            )
    return _vector_store


def _read_file(path: Path, display_name: str) -> list[Document]:
    suffix = path.suffix.lower()
    documents: list[Document] = []
    if suffix == ".txt":
        try:
            text = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            text = path.read_text(encoding="gb18030")
        documents.append(Document(page_content=text, metadata={"source": display_name, "location": "全文"}))
    elif suffix == ".pdf":
        reader = PdfReader(str(path))
        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                documents.append(Document(page_content=text, metadata={"source": display_name, "location": f"第 {index} 页"}))
    elif suffix == ".docx":
        word = WordDocument(str(path))
        paragraphs = [p.text.strip() for p in word.paragraphs if p.text.strip()]
        if paragraphs:
            documents.append(Document(page_content="\n".join(paragraphs), metadata={"source": display_name, "location": "正文"}))
    else:
        raise ValueError("不支持的文件格式")
    if not documents:
        raise ValueError("没有从文件中读取到可用文本；扫描版 PDF 需要先进行 OCR 识别")
    return documents


def index_document(path: Path, document_id: int, display_name: str) -> int:
    documents = _read_file(path, display_name)
    splitter = RecursiveCharacterTextSplitter(chunk_size=700, chunk_overlap=120, separators=["\n\n", "\n", "。", "！", "？", " ", ""])
    chunks = splitter.split_documents(documents)
    for index, chunk in enumerate(chunks, start=1):
        chunk.metadata.update({"document_id": document_id, "chunk_index": index})
    store = _get_store()
    # 重新处理同一份资料时先移除旧片段，避免检索结果重复。
    store.delete(where={"document_id": document_id})
    store.add_documents(chunks, ids=[f"doc-{document_id}-{i}" for i in range(len(chunks))])
    return len(chunks)


async def save_uploaded_file(upload: UploadFile) -> tuple[Path, str, int]:
    suffix = Path(upload.filename or "").suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise HTTPException(400, detail="只支持 PDF、Word（.docx）和 TXT 文件")
    stored_name = f"{uuid.uuid4().hex}{suffix}"
    destination = settings.upload_dir / stored_name
    size = 0
    try:
        with destination.open("wb") as output:
            while content := await upload.read(1024 * 1024):
                size += len(content)
                if size > settings.max_upload_mb * 1024 * 1024:
                    raise HTTPException(400, detail=f"文件不能超过 {settings.max_upload_mb} MB")
                output.write(content)
    except Exception:
        destination.unlink(missing_ok=True)
        raise
    return destination, stored_name, size


def delete_document_vectors(document_id: int) -> None:
    try:
        _get_store().delete(where={"document_id": document_id})
    except Exception:
        logger.exception("删除文档向量失败，document_id=%s", document_id)
        raise


def get_document_chunks(document_id: int) -> list[dict]:
    """Return the exact chunks stored for an administrator to inspect."""
    result = _get_store().get(where={"document_id": document_id}, include=["documents", "metadatas"])
    chunks = []
    for content, metadata in zip(result.get("documents", []), result.get("metadatas", [])):
        chunks.append({
            "index": int(metadata.get("chunk_index", 0)),
            "location": metadata.get("location", "未知位置"),
            "content": content,
        })
    return sorted(chunks, key=lambda item: item["index"])


def _has_keyword_overlap(question: str, content: str) -> bool:
    """Keep Chinese product terms even when compact local vectors have low scores."""
    query_terms = {char for char in question.lower() if char.isalnum() or "\u4e00" <= char <= "\u9fff"}
    content_terms = {char for char in content.lower() if char.isalnum() or "\u4e00" <= char <= "\u9fff"}
    return len(query_terms & content_terms) >= 2


def rebuild_vector_store(documents: list[tuple[int, Path, str]]) -> dict[int, tuple[str, int, str | None]]:
    """清空向量库后重新处理全部资料，返回每个资料的状态。"""
    global _vector_store
    with _load_lock:
        _vector_store = None
        if settings.chroma_dir.exists():
            shutil.rmtree(settings.chroma_dir)
        settings.chroma_dir.mkdir(parents=True, exist_ok=True)
    result = {}
    for document_id, path, name in documents:
        try:
            result[document_id] = ("ready", index_document(path, document_id, name), None)
        except Exception as exc:
            result[document_id] = ("failed", 0, str(exc))
    return result


def ask_knowledge_base(question: str, history: list[tuple[str, str]]) -> tuple[str, list[dict], bool, int]:
    started = time.perf_counter()
    valid = []
    try:
        results = _get_store().similarity_search_with_relevance_scores(question, k=settings.retrieval_k)
        valid = [(doc, score) for doc, score in results if score >= settings.similarity_threshold or _has_keyword_overlap(question, doc.page_content)]
    except Exception as exc:
        # A knowledge-base issue must not remove the model's normal conversation ability.
        logger.warning("知识库检索不可用，将使用模型通用能力回答：%s", exc)
    sources = [{"document": doc.metadata.get("source", "未知文件"), "location": doc.metadata.get("location", "未知位置"), "content": doc.page_content[:500], "score": round(float(score), 3)} for doc, score in valid]
    if not settings.dashscope_api_key:
        raise HTTPException(503, detail="尚未配置大模型 API Key，请检查 backend/.env")
    context = "\n\n".join(f"【资料 {i}｜{doc.metadata.get('source')}｜{doc.metadata.get('location')}】\n{doc.page_content}" for i, (doc, _) in enumerate(valid, start=1)) or "没有检索到相关知识库资料。"
    recent_history = "\n".join(f"{'用户' if role == 'user' else '助手'}：{content}" for role, content in history[-6:]) or "无"
    prompt = ChatPromptTemplate.from_messages([
        ("system", "你是能力完整、乐于助人的通用助手。知识库资料是高价值的补充信息：有相关资料时优先结合它回答；没有相关资料时仍需依靠自己的知识正常回答问候、常识和专业问题，绝不能只回复资料不足。只有当你使用自己的知识补充资料之外的内容时，另起一段并以“【通用补充】”开头。不要编造文件、页码或引用。"),
        ("human", "历史对话：\n{history}\n\n知识库资料：\n{context}\n\n问题：{question}"),
    ])
    try:
        llm = ChatOpenAI(model=settings.llm_model, api_key=settings.dashscope_api_key, base_url=settings.llm_base_url, temperature=0.2)
        answer = (prompt | llm).invoke({"history": recent_history, "context": context, "question": question}).content
    except Exception as exc:
        logger.exception("模型调用失败")
        raise HTTPException(503, detail=f"大模型调用失败，请检查百炼地址、模型名和 API Key：{exc}")
    return str(answer), sources, (not valid) or "【通用补充】" in str(answer), int((time.perf_counter() - started) * 1000)
